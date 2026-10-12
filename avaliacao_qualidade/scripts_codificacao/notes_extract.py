"""Junta, para cada artigo único (CSV), as notas QP da aba FASE 1 e o resumo, por título normalizado."""
import openpyxl, glob, re, csv, json
from collections import defaultdict

root = glob.glob('C:/Repositorios/Survey_VLMs_GSF_PTMS/GSF*.xlsx')[0]
wb = openpyxl.load_workbook(root, read_only=True, data_only=True)
CATS = ['VQA', 'TEXT-IMAGE', 'IMAGE-TEXT', 'RETRIEVAL', 'GENERATE']
norm = lambda s: re.sub(r'[^a-z0-9]', '', str(s).lower())

rows = [tuple(list(r) + [None] * 40)[:40] for r in wb['FASE 1 - SELEÇÃO INICIAL'].iter_rows(values_only=True)]
by_title = defaultdict(list)
cat = None
for r in rows[1:]:
    if isinstance(r[0], str) and r[0].strip().upper() in CATS and all(c is None for c in r[1:28]):
        cat = r[0].strip().upper(); continue
    if r[0] is None:
        continue
    by_title[norm(r[0])].append({'cat': cat, 'abstract': r[10], 'analise': r[28], 'excl': r[29], 'notas': r[30],
                                 'venue': r[3], 'year': r[5], 'doi': r[13]})

f2 = [tuple(list(r) + [None] * 10)[:10] for r in wb['FASE 2 - SELEÇÃO FINAL'].iter_rows(values_only=True)]
print('FASE2 linha exemplo:', [str(c)[:200] for c in f2[2][:4]])

csvrows = list(csv.DictReader(open('C:/Repositorios/Survey_VLMs_GSF_PTMS/avaliacao_qualidade/avaliacao_por_artigo_unico.csv', encoding='utf-8-sig')))
out, miss = {}, 0
for r in csvrows:
    hits = by_title.get(norm(r['Titulo']), [])
    hits = [h for h in hits if not (h['excl'] and str(h['excl']).strip())] or hits
    if not hits:
        miss += 1
    notas = ' || '.join(sorted({str(h['notas']) for h in hits if h['notas']}))
    abstract = next((str(h['abstract']) for h in hits if h['abstract']), '')
    out[r['ID']] = {'title': r['Titulo'], 'cats': r['Categorias'], 'codes': r['QPs_QSs_Fase2'], 'notas': notas,
                    'abstract': abstract, 'fonte': r['Fonte_avaliacao'], 'venue': next((str(h['venue']) for h in hits if h['venue']), '')}
print('sem correspondencia na Fase 1:', miss, 'de', len(csvrows))
print('com notas:', sum(1 for v in out.values() if v['notas']), '| com resumo:', sum(1 for v in out.values() if v['abstract']))
json.dump(out, open('notes606.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
ex = out['P087']
print(ex['notas'][:3000])
